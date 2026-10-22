import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_callout(doc, text, title="NOTE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F0F9FF")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="0284C7"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run_t = p.add_run(f"[{title}] ")
    run_t.bold = True
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = RGBColor(2, 132, 199)
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(15, 23, 42)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

doc = Document()

# Page Margins
sections = doc.sections
for s in sections:
    s.top_margin = Inches(0.8)
    s.bottom_margin = Inches(0.8)
    s.left_margin = Inches(0.8)
    s.right_margin = Inches(0.8)

# ==========================================================
# PAGE 1: COVER PAGE (ATTACHED IMAGE AS FIRST PAPER)
# ==========================================================
cover_p = doc.add_paragraph()
cover_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
cover_run = cover_p.add_run()
if os.path.exists("cover_page.png"):
    cover_run.add_picture("cover_page.png", width=Inches(6.2))
doc.add_page_break()

# Styling functions
def add_section_header(title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor(15, 23, 42) # slate-900

def add_subsection_header(title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(2, 132, 199) # sky-600

def add_body_p(text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.size = Pt(10.5)
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.size = Pt(10)
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    r = p.add_run(text)
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(51, 65, 85)
    return p

# ----------------------------------------------------------
# 1. TEAM DETAILS
# ----------------------------------------------------------
add_section_header("1. Team Details & Academic Affiliation")

tbl_team = doc.add_table(rows=5, cols=3)
tbl_team.alignment = WD_TABLE_ALIGNMENT.CENTER
team_data = [
    ["Parameter", "Details", "Affiliation / Role"],
    ["Student 1", "Sajiya Shaikh (Roll No: 24102A2005)", "Department of Computer Engineering"],
    ["Student 2", "Tamanna Shaikh (Roll No: 24102A2006)", "Department of Computer Engineering"],
    ["Project Supervisor", "Prof. Suja Jayachandran", "Faculty Supervisor, Dept. of Computer Engg."],
    ["Institution & Session", "Vidyalankar Institute of Technology (VIT)", "NAAC Accredited A+ | AY 2026-27 | Subject: NLP"]
]

for row_idx, row in enumerate(tbl_team.rows):
    for col_idx, cell in enumerate(row.cells):
        cell.text = team_data[row_idx][col_idx]
        set_cell_margins(cell, 100, 100, 150, 150)
        if row_idx == 0:
            set_cell_background(cell, "1E3A8A") # Navy
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.size = Pt(10)
        else:
            set_cell_background(cell, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(15, 23, 42)
                    if col_idx == 0:
                        run.font.bold = True

doc.add_paragraph().paragraph_format.space_after = Pt(4)

# ----------------------------------------------------------
# 2. TITLE OF THE PROJECT
# ----------------------------------------------------------
add_section_header("2. Title of the Project")
add_body_p("CYBERBULLYING DETECTION & BIG DATA ANALYTICS SYSTEM", bold_prefix="Project Title: ")
add_body_p("An end-to-end intelligent NLP and distributed Big Data framework designed for automated detection, fine-grained toxicity categorization, real-time moderation escalation, and longitudinal analytics of online social commentary.", bold_prefix="Subtitle & Scope: ")

# ----------------------------------------------------------
# 3. DETAILED PROBLEM STATEMENT
# ----------------------------------------------------------
add_section_header("3. Detailed Problem Statement")
add_body_p("The exponential proliferation of social media platforms, discussion forums, and collaborative networks has introduced critical digital safety challenges. Rampant online toxicity, cyberstalking, hate speech, identity-based harassment, and cyberbullying severely affect user well-being and community health. Conventional platform moderation mechanisms suffer from severe foundational bottlenecks:")

add_bullet("Human moderators cannot manually inspect millions of incoming comments daily, leading to prohibitive latency and emotional fatigue.", "Operational Scalability Bottleneck: ")
add_bullet("Cyberbullying frequently leverages obfuscated vocabulary, colloquialisms, sarcastic framing, and intentional character misspellings (e.g., 'f*ck', 'loser', 'fugly') that evade elementary keyword blacklists.", "Linguistic Nuance & Adversarial Text: ")
add_bullet("Severe class imbalance (where toxic comments constitute 5-10% of total social commentary) causes models to suffer high false alarm rates or extreme false negative omissions.", "Toxicity Class Imbalance: ")
add_bullet("Traditional architectures fail to bridge the operational gap between localized real-time inference (sub-50ms web scoring) and macroscopic longitudinal analytics over millions of historical comments.", "Absence of Unified Big Data Synergy: ")

add_callout(doc, "Project Objective: To build a robust Natural Language Processing classification pipeline combined with an Apache PySpark distributed Big Data processing engine and an interactive Flask moderation portal to achieve zero-leakage, high-precision detection and multi-million-row analytics.", "CORE OBJECTIVE")

# ----------------------------------------------------------
# 4. TECHNOLOGY / ALGORITHM USED
# ----------------------------------------------------------
add_section_header("4. Technology / Algorithm Used")
add_subsection_header("4.1 Machine Learning & Natural Language Processing Pipeline")

add_bullet("Lowercasing, punctuation stripping, regex noise suppression, and domain-tailored stopword elimination.", "Text Preprocessing & Cleaning: ")
add_bullet("A hybrid scikit-learn FeatureUnion combining Word N-Grams (range 1-2, sublinear TF-IDF scaling) and Character N-Grams (range 3-5). Character n-grams provide crucial resilience against adversarial typos and masked expletives.", "Feature Engineering (FeatureUnion TF-IDF): ")
add_bullet("Linear classifier with hinge loss and elastic-net/L2 regularization. Best-performing model achieving 92.44% accuracy and 92.48% F1-score with ultra-low latency.", "SGDClassifier (Stochastic Gradient Descent / Linear SVM): ")
add_bullet("Evaluated with calibrated probabilities, yielding 91.89% accuracy and an outstanding 97.10% ROC-AUC for continuous risk ranking.", "Logistic Regression (Calibrated Probability): ")
add_bullet("LinearSVC, Decision Tree Classifier (91.86% acc), Multinomial Naive Bayes (86.94% acc), and AdaBoost (80.05% acc) for rigorous comparative benchmarking.", "Ensemble & Baseline Algorithms: ")

add_subsection_header("4.2 Distributed Big Data Engineering & Analytics")
add_bullet("Apache PySpark 4.2.0 on Java 21 LTS runtime for parallelized processing of 1.8M real-world toxic and non-toxic comments.", "Distributed Compute Engine: ")
add_bullet("Partitioned columnar Snappy-compressed Parquet storage achieving a 94.6% disk reduction (41.8 MB vs 778.4 MB) and 17.9x query speedup.", "Storage Layer (Snappy Parquet): ")
add_bullet("Dense ranking of high-risk sources (dense_rank) and cumulative longitudinal toxicity growth over time (unboundedPreceding window calculations).", "Spark SQL & Window Functions: ")

add_subsection_header("4.3 Web Platform & Content Moderation Architecture")
add_bullet("Python Flask 3.1.3 microservices backend with modular Jinja2 templating.", "Web Framework: ")
add_bullet("Thread-safe SQLite platform database logging all analyzed cases, risk percentages, severity tiers, and moderation actions.", "Database & Audit Store: ")
add_bullet("Real-time manual text analysis, multi-case live comment stream simulation, and administrative moderation dashboard.", "Interactive Web Interface: ")

# Insert Architecture Figure
if os.path.exists("report_assets/fig4_architecture_pipeline.png"):
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(2)
    p_img.add_run().add_picture("report_assets/fig4_architecture_pipeline.png", width=Inches(6.2))
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cap = p_cap.add_run("Figure 1: End-to-End System Architecture — Data Ingestion, NLP Pipeline, Distributed PySpark Analytics, and Moderation Web Platform")
    r_cap.font.size = Pt(8.5)
    r_cap.font.italic = True
    r_cap.font.color.rgb = RGBColor(100, 116, 139)

# ----------------------------------------------------------
# 5. DATASET DESCRIPTION
# ----------------------------------------------------------
add_section_header("5. Dataset Description")
add_body_p("The system is engineered and validated across two complementary tiers of datasets: a clean domain-specific cyberbullying ground-truth dataset and an enterprise-scale Big Data toxicity corpus.")

tbl_data = doc.add_table(rows=4, cols=4)
tbl_data.alignment = WD_TABLE_ALIGNMENT.CENTER
ds_content = [
    ["Dataset Name", "Scale & File Size", "Target Labels / Features", "Operational Role in System"],
    ["Core Cyberbullying Corpus\n(dataset.csv)", "17,065 Clean Rows\n4.91 MB", "Binary: Bullying (-1) [62.0%]\nNon-Bullying (0) [38.0%]", "Primary training & clean held-out zero-leakage evaluation partition."],
    ["Civil Comments Big Data\n(civil_comments_train.csv)", "1,804,874 Rows\n778.4 MB (CSV) / 41.8 MB (Parquet)", "Continuous Toxicity [0,1], Insult, Threat, Obscene, Identity Attack", "Big Data analytics, PySpark window functions, and cross-domain data enrichment."],
    ["Instagram Comments Dataset\n(instagram_comments.csv)", "82.75 MB Raw Text", "Unstructured Social Comments, Platform User Feedback", "Live simulation validation and colloquial text robustness testing."]
]

for r_idx, row in enumerate(tbl_data.rows):
    for c_idx, cell in enumerate(row.cells):
        cell.text = ds_content[r_idx][c_idx]
        set_cell_margins(cell, 80, 80, 120, 120)
        if r_idx == 0:
            set_cell_background(cell, "0284C7") # Sky
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.size = Pt(9.5)
        else:
            set_cell_background(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(15, 23, 42)
                    if c_idx == 0:
                        run.font.bold = True

add_callout(doc, "Zero Data Leakage Rigor: The core dataset was deduplicated using case-folded normalization, conflicting annotations were deterministically eliminated, and an 80/20 stratified split was enforced with 0.00% train/test duplicate overlap.", "DATA INTEGRITY")

# ----------------------------------------------------------
# 6. SCREENSHOT / RESULT
# ----------------------------------------------------------
add_section_header("6. Screenshots & Experimental Results")
add_subsection_header("6.1 Model Performance Benchmarks on Held-Out Test Set (N=3,413)")

tbl_res = doc.add_table(rows=7, cols=6)
tbl_res.alignment = WD_TABLE_ALIGNMENT.CENTER
bench_data = [
    ["Algorithm / Model", "Accuracy (%)", "Precision (%)", "Recall (%)", "F1-Score (%)", "ROC-AUC (%)"],
    ["SGDClassifier (Linear SVM)", "92.44%", "92.63%", "92.44%", "92.48%", "N/A (Hinge)"],
    ["Logistic Regression", "91.89%", "92.05%", "91.89%", "91.93%", "97.10%"],
    ["Decision Tree Classifier", "91.86%", "91.87%", "91.86%", "91.86%", "91.75%"],
    ["Linear SVC", "91.33%", "91.48%", "91.33%", "91.37%", "N/A"],
    ["Multinomial Naive Bayes", "86.94%", "87.29%", "86.94%", "86.61%", "95.01%"],
    ["AdaBoost Classifier", "80.05%", "85.60%", "80.05%", "80.28%", "87.59%"]
]

for r_idx, row in enumerate(tbl_res.rows):
    for c_idx, cell in enumerate(row.cells):
        cell.text = bench_data[r_idx][c_idx]
        set_cell_margins(cell, 70, 70, 100, 100)
        if r_idx == 0:
            set_cell_background(cell, "1E293B") # Dark slate
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.size = Pt(9.5)
        else:
            is_best = (r_idx == 1)
            set_cell_background(cell, "ECFDF5" if is_best else ("F8FAFC" if r_idx % 2 == 1 else "FFFFFF"))
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(6, 95, 70) if is_best else RGBColor(15, 23, 42)
                    if c_idx == 0 or is_best:
                        run.font.bold = True

# Insert Performance Graph
if os.path.exists("report_assets/fig1_model_performance.png"):
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(2)
    p_img.add_run().add_picture("report_assets/fig1_model_performance.png", width=Inches(6.0))
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cap = p_cap.add_run("Figure 2: Empirical Performance Benchmark across Six NLP Classifiers on Unseen Test Partition")
    r_cap.font.size = Pt(8.5)
    r_cap.font.italic = True
    r_cap.font.color.rgb = RGBColor(100, 116, 139)

# Insert Confusion Matrix
if os.path.exists("report_assets/fig2_confusion_matrices.png"):
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(2)
    p_img.add_run().add_picture("report_assets/fig2_confusion_matrices.png", width=Inches(6.0))
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cap = p_cap.add_run("Figure 3: Confusion Matrix Evaluation for SGDClassifier (Best Overall) and Logistic Regression (Best Probability Calibration)")
    r_cap.font.size = Pt(8.5)
    r_cap.font.italic = True
    r_cap.font.color.rgb = RGBColor(100, 116, 139)

add_subsection_header("6.2 Big Data PySpark Distributed Analytics Results (1.8M Records)")
add_body_p("The distributed PySpark pipeline processed 1,804,874 real comments with zero data corruption. Columnar Snappy-compressed Parquet reduced physical storage footprint from 778.4 MB down to 41.8 MB (a 94.6% space reduction). Query read latency dropped from 18.17 seconds to 1.015 seconds (17.9x acceleration), enabling interactive analytical response times.")

# Insert Big Data Chart
if os.path.exists("report_assets/fig3_bigdata_analytics.png"):
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(2)
    p_img.add_run().add_picture("report_assets/fig3_bigdata_analytics.png", width=Inches(6.1))
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cap = p_cap.add_run("Figure 4: Big Data PySpark Engineering Benchmarks — Class Distribution, Subtype Prevalence, Parquet Compression Ratio, and Read Latency Speedup")
    r_cap.font.size = Pt(8.5)
    r_cap.font.italic = True
    r_cap.font.color.rgb = RGBColor(100, 116, 139)

add_subsection_header("6.3 Web Application Interface & Content Moderation Showcase")
add_body_p("The production web application was deployed on Flask 3.1.3, providing immediate content analysis, risk scores (85% High Severity vs 10% Low Severity), automated database audit persistence in SQLite, and active moderation management.")

# Insert Web App Graphic
if os.path.exists("report_assets/fig5_web_interface_simulation.png"):
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(2)
    p_img.add_run().add_picture("report_assets/fig5_web_interface_simulation.png", width=Inches(6.1))
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cap = p_cap.add_run("Figure 5: Production Web Dashboard Interface — Real-Time Text Analysis, Immediate Risk Tagging, and Moderation Escalation Stream")
    r_cap.font.size = Pt(8.5)
    r_cap.font.italic = True
    r_cap.font.color.rgb = RGBColor(100, 116, 139)

# ----------------------------------------------------------
# 7. CONCLUSION
# ----------------------------------------------------------
add_section_header("7. Conclusion & Future Work")
add_body_p("The Cyberbullying Detection & Big Data Analytics System successfully demonstrates an enterprise-ready, dual-tier solution addressing toxic online commentary. By uniting robust Natural Language Processing classification with distributed Apache PySpark data engineering, the system achieves remarkable accuracy while seamlessly scaling to millions of records.")

add_bullet("The SGDClassifier linear SVM achieved 92.44% accuracy, 92.63% precision, and 92.48% F1-score on a strictly held-out zero-leakage test set, proving highly resilient to adversarial typos via character n-grams.", "NLP Classification Efficacy: ")
add_bullet("Snappy Parquet columnar encoding reduced storage by 94.6% and delivered an 17.9x query acceleration over raw CSVs, demonstrating high-efficiency handling of 1.8M comments.", "Big Data Throughput: ")
add_bullet("The Flask web application provides sub-50ms inference, interactive moderation escalation, and automated persistent auditing in SQLite.", "Production Deployment: ")
add_bullet("Integration of contextual transformer models (e.g., RoBERTa/DeBERTa), multilingual cross-lingual toxicity detection, and automated user feedback loops.", "Future Scope: ")

# ----------------------------------------------------------
# 8. GITHUB LINK OF PROJECT CODE
# ----------------------------------------------------------
add_section_header("8. GitHub Project Repository & Code Link")
add_body_p("The complete source code, trained machine learning model artifacts, PySpark distributed analytics scripts, datasets, and Flask web application are publicly hosted and version-controlled at:")

p_git = doc.add_paragraph()
p_git.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_git.paragraph_format.space_before = Pt(6)
p_git.paragraph_format.space_after = Pt(12)
r_git = p_git.add_run("https://github.com/Shaziya-1/Cyber-Bullying-Detection-System")
r_git.bold = True
r_git.font.size = Pt(13)
r_git.font.color.rgb = RGBColor(2, 132, 199)
r_git.font.underline = True

tbl_repo = doc.add_table(rows=4, cols=2)
tbl_repo.alignment = WD_TABLE_ALIGNMENT.CENTER
repo_details = [
    ["Repository Parameter", "Value"],
    ["GitHub Repository URL", "https://github.com/Shaziya-1/Cyber-Bullying-Detection-System"],
    ["Main Application Entry Point", "app.py (Flask Web Application & Moderation Service)"],
    ["Core Deliverables", "ML Models (.pkl), Big Data PySpark Analytics, SQL Moderation DB, Datasets"]
]

for r_idx, row in enumerate(tbl_repo.rows):
    for c_idx, cell in enumerate(row.cells):
        cell.text = repo_details[r_idx][c_idx]
        set_cell_margins(cell, 80, 80, 120, 120)
        if r_idx == 0:
            set_cell_background(cell, "1E3A8A")
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.size = Pt(9.5)
        else:
            set_cell_background(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    if c_idx == 0:
                        run.font.bold = True

doc.save("Mini_Project_Report_Cyberbullying_Detection.docx")
print("Saved DOCX successfully as Mini_Project_Report_Cyberbullying_Detection.docx")
